"""Build the planning workbook from the demo database.

    python planning/build_workbook.py

Writes planning/store_planning_model.xlsx: new store unit economics, a 3-year
plan by channel, and pricing and promotion tests. The last 30 days of actuals
are pasted as values; every other number is an Excel formula, so changing an
input recalculates the whole model. Before writing, the actuals are tied out
a second way (stores vs categories, item lines vs POS transactions), the same
idea as the Copilot's verifier.

Excel calculates the file on open. The committed copy was also recalculated
with LibreOffice so file previews show numbers.
"""
from __future__ import annotations

import sqlite3
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from data.generate import DB_PATH, ensure_db, load_config  # noqa: E402

OUT = ROOT / "planning" / "store_planning_model.xlsx"
DAYS = 30
MONTHS = 60
YEARS = ("2027", "2028", "2029")
SCENARIOS = ("Base", "Upside", "Downside")
CAFE = "Cafe & Hot Bar"
UNMAPPED = "Unmapped POS codes"
DEFAULT_CATEGORY = "Meat & Seafood"

# ----------------------------------------------------------------------------- styles
BLUE, BLACK, GREEN, GRAY, WHITE, ACCENT = "FF0000FF", "FF000000", "FF008000", "FF6B6B6B", "FFFFFFFF", "FF2E5E45"
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
BAND = PatternFill("solid", fgColor=ACCENT)
LINE = Side(style="thin", color="FFBFBFBF")
USD = '$#,##0;[Red]($#,##0);"-"'
USD_K = '$#,##0,"K";[Red]($#,##0,"K");"-"'
USD2 = '$#,##0.00;[Red]($#,##0.00);"-"'
PCT = '0.0%;[Red](0.0%);"-"'
PCT0 = '0%;[Red](0%);"-"'
PTS = '+0.0%;[Red]-0.0%;0.0%'
NUM = '#,##0;[Red](#,##0);"-"'
DEC = '0.0'
COLORS = {"in": BLUE, "f": BLACK, "link": GREEN, "text": BLACK, "note": GRAY}


def put(ws, ref, value, kind="f", fmt=None, bold=False, fill=None, wrap=False, align=None, size=10):
    c = ws[ref]
    c.value = value
    c.font = Font(name="Arial", size=size, bold=bold, color=COLORS[kind], italic=False)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    indent = 1 if kind == "note" and not ref.startswith("A") else 0
    if wrap or align or indent:
        c.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="top" if wrap else None, indent=indent)
    return c


def title(ws, text, sub):
    put(ws, "A1", text, "text", bold=True, size=14).font = Font(name="Arial", size=14, bold=True, color=ACCENT)
    put(ws, "A2", sub, "note")
    ws.sheet_view.showGridLines = False


def section(ws, row, text, last_col):
    for c in range(1, last_col + 1):
        ws.cell(row, c).fill = BAND
    put(ws, f"A{row}", text, "text", bold=True).font = Font(name="Arial", size=10, bold=True, color=WHITE)


def headers(ws, row, labels, start=1):
    for i, label in enumerate(labels):
        c = put(ws, f"{get_column_letter(start + i)}{row}", label, "text", bold=True, wrap=True,
                align="left" if i == 0 or label in ("Unit", "Notes", "Channel") else "right")
        c.border = Border(bottom=LINE)
        if label == "Notes":
            c.alignment = Alignment(wrap_text=True, horizontal="left", vertical="top", indent=1)


def widths(ws, spec):
    for letter, w in spec.items():
        ws.column_dimensions[letter].width = w


def quote(sheet):
    return f"'{sheet}'" if any(ch in sheet for ch in " -&") else sheet


def absref(sheet, ref):
    letters = "".join(ch for ch in ref if ch.isalpha())
    digits = ref[len(letters):]
    return f"{quote(sheet)}!${letters}${digits}"


def name(wb, nm, sheet, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=absref(sheet, ref))


def note(ws, ref, text):
    ws[ref].comment = Comment(text, "Store Copilot")


# ----------------------------------------------------------------------------- data
def load_actuals(db: Path) -> dict:
    cfg = load_config()
    since = f"date('{cfg['data_end_date']}', '-{DAYS} day')"
    con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    try:
        q = lambda sql: con.execute(sql).fetchall()  # noqa: E731
        stores = q("SELECT store_id, name, neighborhood FROM stores ORDER BY store_id")
        items = {r[0]: r[1:] for r in q(f"""SELECT store_id, SUM(revenue), SUM(cogs) FROM sales_daily
                                            WHERE sale_date > {since} GROUP BY store_id""")}
        tx = {r[0]: r[1:] for r in q(f"""SELECT store_id, COUNT(*), SUM(revenue), SUM(transactions),
                                                SUM(member_transactions)
                                         FROM transactions_daily WHERE sale_date > {since} GROUP BY store_id""")}
        waste = dict(q(f"SELECT store_id, SUM(waste_cost) FROM waste_daily WHERE waste_date > {since} GROUP BY store_id"))
        targets = q("SELECT category, target_margin_pct / 100.0 FROM category_targets ORDER BY rowid")
        cats = {r[0]: r[1:] for r in q(f"""SELECT COALESCE(m.category, '{UNMAPPED}'), SUM(sd.revenue), SUM(sd.cogs),
                                                  SUM(sd.units)
                                           FROM sales_daily sd LEFT JOIN sku_master m
                                             ON m.master_sku_id = sd.master_sku_id
                                           WHERE sd.sale_date > {since} GROUP BY 1""")}
    finally:
        con.close()

    store_rows = []
    for sid, sname, hood in stores:
        days, tx_rev, n_tx, n_member = tx[sid]
        rev, cogs = items[sid]
        label = f"{sname} ({hood.replace(' (hypothetical)', ', hypothetical')})"
        store_rows.append(dict(id=sid, label=label, days=days, rev=rev, tx_rev=tx_rev, cogs=cogs, tx=n_tx,
                               member_tx=n_member, waste=waste.get(sid, 0.0)))
    cat_rows = [dict(name=c, target=t, rev=cats[c][0], cogs=cats[c][1], units=cats[c][2]) for c, t in targets]
    if UNMAPPED in cats:
        cat_rows.append(dict(name=UNMAPPED, target=None, rev=cats[UNMAPPED][0], cogs=cats[UNMAPPED][1],
                             units=cats[UNMAPPED][2]))

    # Tie out before writing anything: three routes to the same totals.
    by_store = sum(s["rev"] for s in store_rows)
    for label, other in (("POS transactions", sum(s["tx_rev"] for s in store_rows)),
                         ("categories", sum(c["rev"] for c in cat_rows))):
        if abs(by_store - other) > 0.005 * by_store:
            raise SystemExit(f"Actuals don't tie: stores {by_store:,.0f} vs {label} {other:,.0f}")
    cogs_store, cogs_cat = sum(s["cogs"] for s in store_rows), sum(c["cogs"] for c in cat_rows)
    if abs(cogs_store - cogs_cat) > 0.005 * cogs_store:
        raise SystemExit(f"Cost of goods doesn't tie: stores {cogs_store:,.0f} vs categories {cogs_cat:,.0f}")
    return dict(stores=store_rows, cats=cat_rows, end=str(cfg["data_end_date"]))


# ----------------------------------------------------------------------------- sheets
def build_actuals(wb, ws, a):
    end = date.fromisoformat(a["end"])
    start = end - timedelta(days=DAYS - 1)
    title(ws, "Actuals", f"Last {DAYS} days of the demo data, {start} to {end}. Synthetic. Pasted from store.db by "
                         "planning/build_workbook.py. Blue = pasted value, black = formula.")
    put(ws, "A3", "Days in window", "text")
    put(ws, "B3", DAYS, "in", NUM)
    name(wb, "days_window", ws.title, "B3")

    section(ws, 5, "By store", 14)
    headers(ws, 6, ["Store", "Days open", "Sales, item lines", "Sales, POS transactions", "Cost of goods",
                    "Gross margin", "Transactions", "Member transactions", "Average basket", "Member share",
                    "Waste cost", "Waste, % of sales", "Sales per open day", "Annual run-rate"])
    r0 = 7
    for i, s in enumerate(a["stores"]):
        r = r0 + i
        put(ws, f"A{r}", s["label"], "text")
        put(ws, f"B{r}", s["days"], "in", NUM)
        put(ws, f"C{r}", round(s["rev"], 2), "in", USD)
        put(ws, f"D{r}", round(s["tx_rev"], 2), "in", USD)
        put(ws, f"E{r}", round(s["cogs"], 2), "in", USD)
        put(ws, f"F{r}", f"=IF(C{r}=0,0,(C{r}-E{r})/C{r})", "f", PCT)
        put(ws, f"G{r}", s["tx"], "in", NUM)
        put(ws, f"H{r}", s["member_tx"], "in", NUM)
        put(ws, f"I{r}", f"=IF(G{r}=0,0,C{r}/G{r})", "f", USD2)
        put(ws, f"J{r}", f"=IF(G{r}=0,0,H{r}/G{r})", "f", PCT)
        put(ws, f"K{r}", round(s["waste"], 2), "in", USD)
        put(ws, f"L{r}", f"=IF(C{r}=0,0,K{r}/C{r})", "f", PCT)
        put(ws, f"M{r}", f"=IF(B{r}=0,0,C{r}/B{r})", "f", USD)
        put(ws, f"N{r}", f"=M{r}*365", "f", USD)
    rt = r0 + len(a["stores"])
    put(ws, f"A{rt}", "Total", "text", bold=True)
    for c in "CDEGHKMN":
        put(ws, f"{c}{rt}", f"=SUM({c}{r0}:{c}{rt - 1})", "f", USD if c in "CDEKMN" else NUM, bold=True)
    put(ws, f"F{rt}", f"=IF(C{rt}=0,0,(C{rt}-E{rt})/C{rt})", "f", PCT, bold=True)
    put(ws, f"I{rt}", f"=IF(G{rt}=0,0,C{rt}/G{rt})", "f", USD2, bold=True)
    put(ws, f"J{rt}", f"=IF(G{rt}=0,0,H{rt}/G{rt})", "f", PCT, bold=True)
    put(ws, f"L{rt}", f"=IF(C{rt}=0,0,K{rt}/C{rt})", "f", PCT, bold=True)
    for c in range(1, 15):
        ws.cell(rt, c).border = Border(top=LINE)
    s3 = a["stores"][-1]
    put(ws, f"A{rt + 1}", f"{s3['label'].split(' (')[0]} opened recently: {s3['days']} days in the window, "
                          "so its run-rate understates a full year.", "note")
    ids = {s["id"]: r0 + i for i, s in enumerate(a["stores"])}
    name(wb, "s1_per_day", ws.title, f"M{ids['S1']}")
    name(wb, "s3_per_day", ws.title, f"M{ids['S3']}")
    name(wb, "flag_runrate", ws.title, f"N{ids['S1']}")
    name(wb, "stores_runrate", ws.title, f"N{rt}")
    name(wb, "gm_actual", ws.title, f"F{rt}")
    name(wb, "waste_actual", ws.title, f"L{rt}")
    name(wb, "member_tx", ws.title, f"H{rt}")

    cs = rt + 3
    section(ws, cs, "By category", 14)
    headers(ws, cs + 1, ["Category", "Channel", "Sales", "Cost of goods", "Units", "Gross margin", "Target margin",
                         "Gap to target", "Average price", "Average unit cost", "Units per week", "Share of sales"])
    c0 = cs + 2
    for i, c in enumerate(a["cats"]):
        r = c0 + i
        put(ws, f"A{r}", c["name"], "text")
        put(ws, f"B{r}", "Cafe & hot bar" if c["name"] == CAFE else "Grocery", "text")
        put(ws, f"C{r}", round(c["rev"], 2), "in", USD)
        put(ws, f"D{r}", round(c["cogs"], 2), "in", USD)
        put(ws, f"E{r}", c["units"], "in", NUM)
        put(ws, f"F{r}", f"=IF(C{r}=0,0,(C{r}-D{r})/C{r})", "f", PCT)
        if c["target"] is not None:
            put(ws, f"G{r}", c["target"], "in", PCT)
            put(ws, f"H{r}", f"=F{r}-G{r}", "f", PTS)
        put(ws, f"I{r}", f"=IF(E{r}=0,0,C{r}/E{r})", "f", USD2)
        put(ws, f"J{r}", f"=IF(E{r}=0,0,D{r}/E{r})", "f", USD2)
        put(ws, f"K{r}", f"=E{r}/days_window*7", "f", NUM)
    ct = c0 + len(a["cats"])
    for i in range(len(a["cats"])):
        put(ws, f"L{c0 + i}", f"=IF($C${ct}=0,0,C{c0 + i}/$C${ct})", "f", PCT)
    put(ws, f"A{ct}", "Total", "text", bold=True)
    for c in "CDE":
        put(ws, f"{c}{ct}", f"=SUM({c}{c0}:{c}{ct - 1})", "f", USD if c != "E" else NUM, bold=True)
    put(ws, f"F{ct}", f"=IF(C{ct}=0,0,(C{ct}-D{ct})/C{ct})", "f", PCT, bold=True)
    for c in range(1, 15):
        ws.cell(ct, c).border = Border(top=LINE)
    if a["cats"][-1]["name"] == UNMAPPED:
        put(ws, f"A{ct + 1}", "Unmapped POS codes are sales the POS can't tie to a product; counted as grocery.", "note")
    n_named = sum(1 for c in a["cats"] if c["target"] is not None)
    wb.defined_names["category_list"] = DefinedName("category_list",
                                                    attr_text=f"{quote(ws.title)}!$A${c0}:$A${c0 + n_named - 1}")

    ch = ct + 3
    section(ws, ch, "By channel", 14)
    headers(ws, ch + 1, ["Channel", "Sales", "Cost of goods", "Gross margin", "Share of sales", "Annual run-rate"])
    rng = f"$B${c0}:$B${ct - 1}"
    for i, label in enumerate(("Grocery", "Cafe & hot bar")):
        r = ch + 2 + i
        put(ws, f"A{r}", label, "text")
        put(ws, f"B{r}", f'=SUMIF({rng},"{label}",$C${c0}:$C${ct - 1})', "f", USD)
        put(ws, f"C{r}", f'=SUMIF({rng},"{label}",$D${c0}:$D${ct - 1})', "f", USD)
        put(ws, f"D{r}", f"=IF(B{r}=0,0,(B{r}-C{r})/B{r})", "f", PCT)
        put(ws, f"E{r}", f"=IF($C${ct}=0,0,B{r}/$C${ct})", "f", PCT)
        put(ws, f"F{r}", f"=E{r}*stores_runrate", "f", USD)
    put(ws, f"A{ch + 4}", "Run-rate splits the stores' annual run-rate by each channel's share of sales.", "note")
    name(wb, "grocery_gm", ws.title, f"D{ch + 2}")
    name(wb, "cafe_gm", ws.title, f"D{ch + 3}")
    name(wb, "grocery_runrate", ws.title, f"F{ch + 2}")
    name(wb, "cafe_runrate", ws.title, f"F{ch + 3}")

    widths(ws, {"A": 40, **{get_column_letter(i): 14 for i in range(2, 15)}})
    ws.row_dimensions[6].height = 30
    ws.row_dimensions[cs + 1].height = 30
    return dict(store_total=rt, store_first=r0, cat_first=c0, cat_total=ct, s3_days=s3["days"],
                s3_name=s3["label"].split(" (")[0])


# (name, label, unit, (base, upside, downside), format, note)
SCENARIO_INPUTS = [
    ("New store", [
        ("prod", "Sales per sq ft vs flagship", "% of flagship", (1.00, 1.10, 0.85), PCT0,
         "100% means the same sales per sq ft as the flagship."),
        ("ramp_start", "Month 1 sales", "% of mature", (0.60, 0.70, 0.50), PCT0, None),
        ("ramp_months", "Months to reach mature sales", "months", (18, 12, 24), NUM,
         "The demo data ramps a store in 30 days, which is fast. Real stores take 12 to 24 months."),
        ("gm_delta", "Gross margin change vs last 30 days", "points", (0.0, 0.01, -0.015), PTS,
         "Added to each channel's actual margin. The Last 30 days column shows the blended actual."),
        ("waste_pct", "Waste", "% of store sales", (None, 0.016, 0.025), PCT,
         "Base is the actual. The Perishables SOP aims to cut waste."),
        ("labor_pct", "Store labor", "% of sales", (0.19, 0.18, 0.21), PCT,
         "Not in the data. The cafe and hot bar make this higher than a typical grocery store."),
        ("labor_floor", "Labor floor while ramping", "% of mature labor", (0.85, 0.80, 0.90), PCT0,
         "A new store can't cut staff in line with lower early sales."),
        ("rent_psf", "Rent", "$ per sq ft per year", (60, 55, 70), USD, "Not in the data. Assumption."),
        ("other_pct", "Other operating costs", "% of sales", (0.08, 0.075, 0.09), PCT,
         "Card fees, utilities, supplies, repairs, local marketing."),
        ("buildout_psf", "Build-out", "$ per sq ft", (250, 225, 300), USD, "Construction, equipment, fixtures."),
        ("preopen", "Pre-opening costs", "$", (150000, 125000, 200000), USD, "Hiring, training, launch marketing."),
    ]),
    ("Growth", [
        ("g_grocery", "Grocery sales growth, existing stores", "% per year", (0.04, 0.07, 0.01), PCT, None),
        ("g_cafe", "Cafe and hot bar growth, existing stores", "% per year", (0.08, 0.12, 0.03), PCT, None),
        ("open_2027", "New stores opened in 2027", "stores", (1, 1, 0), NUM,
         "Opened at the start of the year; each follows the New Store sheet."),
        ("open_2028", "New stores opened in 2028", "stores", (1, 2, 1), NUM, None),
        ("open_2029", "New stores opened in 2029", "stores", (1, 2, 0), NUM, None),
        ("cat_y1", "Catering sales in 2027", "$", (250000, 400000, 100000), USD, "New channel. Not in the data."),
        ("g_cat", "Catering growth", "% per year", (0.30, 0.50, 0.10), PCT, None),
        ("cat_margin", "Catering contribution margin", "% after food and labor", (0.25, 0.30, 0.20), PCT, None),
        ("app_pct", "App orders: extra sales by 2029", "% of store sales", (0.04, 0.06, 0.02), PCT,
         "Only sales the app adds, not orders that move from the store to the app."),
        ("app_cost", "App fulfillment cost", "% of app sales", (0.12, 0.10, 0.15), PCT,
         "Picking, packing, delivery subsidy."),
    ]),
    ("Membership", [
        ("g_members", "Member growth", "% per year", (0.20, 0.35, 0.10), PCT, None),
        ("paid_conv", "Members who pay a fee", "% of members", (0.25, 0.35, 0.15), PCT, None),
        ("perks_cost", "Member perks cost", "% of fees", (0.30, 0.25, 0.40), PCT, None),
    ]),
    ("Company", [
        ("ho_pct", "Head office costs", "% of revenue", (0.05, 0.045, 0.06), PCT,
         "Finance, buying, marketing, leadership. Not in the data."),
    ]),
]

# (name, label, unit, value, format, note)
FIXED_INPUTS = [
    ("sqft_s1", "Flagship size", "sq ft", 12000, NUM, "Store sizes are not in the data. Assumption."),
    ("sqft_s2", "Store 2 size", "sq ft", 8500, NUM, None),
    ("sqft_s3", "Store 3 size", "sq ft", 6500, NUM, None),
    ("sqft_new", "New store size", "sq ft", 8000, NUM, None),
    ("inventory", "Opening inventory", "$", 200000, USD, None),
    ("visits", "Member visits per month", "visits", 4, NUM, "Turns member transactions into a member count."),
    ("fee", "Annual membership fee", "$", 99, USD, "Set to 0 to model a free program."),
    ("app_ramp_27", "App ramp, 2027", "% of 2029 level", 1 / 3, PCT0, None),
    ("app_ramp_28", "App ramp, 2028", "% of 2029 level", 2 / 3, PCT0, None),
    ("app_ramp_29", "App ramp, 2029", "% of 2029 level", 1.0, PCT0, None),
]


def build_inputs(wb, ws, act):
    title(ws, "Inputs", "Change the blue cells. Pick a scenario in C4; the Live column feeds every other sheet.")
    put(ws, "A4", "Scenario", "text", bold=True)
    put(ws, "C4", "Base", "in", fill=YELLOW, bold=True)
    dv = DataValidation(type="list", formula1='"Base,Upside,Downside"', allow_blank=False)
    dv.error, dv.errorTitle = "Pick Base, Upside or Downside.", "Scenario"
    ws.add_data_validation(dv)
    dv.add("C4")
    name(wb, "scenario", ws.title, "C4")
    put(ws, "D4", "Pick Base, Upside or Downside from the list.", "note")

    headers(ws, 6, ["Assumption", "Unit", "Live", *SCENARIOS, "Last 30 days", "Notes"])
    actual_col = {
        "gm_delta": "=gm_actual",
        "waste_pct": "=waste_actual",
        "ramp_start": "=IF(s1_per_day=0,0,(s3_per_day/sqft_s3)/(s1_per_day/sqft_s1))",
        "prod": None,
    }
    r = 7
    for sec, items in SCENARIO_INPUTS:
        section(ws, r, sec, 8)
        r += 1
        for key, label, unit, values, fmt, text in items:
            put(ws, f"A{r}", label, "text")
            put(ws, f"B{r}", unit, "note")
            put(ws, f"C{r}", f"=INDEX(D{r}:F{r},MATCH(scenario,$D$6:$F$6,0))", "f", fmt, bold=True)
            for j, v in enumerate(values):
                ref = f"{'DEF'[j]}{r}"
                if v is None:  # Base waste is the actual, rounded to 0.1%
                    put(ws, ref, "=ROUND(waste_actual,3)", "link", fmt)
                else:
                    put(ws, ref, v, "in", fmt)
            if actual_col.get(key):
                put(ws, f"G{r}", actual_col[key], "link", PCT)
            if key == "ramp_start":
                text = (f"Last 30 days column: {act['s3_name']}'s sales per sq ft in its first {act['s3_days']} "
                        "days, vs the flagship.")
            if text:
                put(ws, f"H{r}", text, "note")
            name(wb, key, ws.title, f"C{r}")
            name(wb, f"{key}_base", ws.title, f"D{r}")
            r += 1
        r += 1
    section(ws, r, "Same in every scenario", 8)
    r += 1
    for key, label, unit, value, fmt, text in FIXED_INPUTS:
        put(ws, f"A{r}", label, "text")
        put(ws, f"B{r}", unit, "note")
        put(ws, f"C{r}", value, "in", fmt, bold=True)
        if text:
            put(ws, f"H{r}", text, "note")
        name(wb, key, ws.title, f"C{r}")
        r += 1
    widths(ws, {"A": 40, "B": 20, "C": 11, "D": 11, "E": 11, "F": 11, "G": 12, "H": 80})
    ws.freeze_panes = "C7"


def build_new_store(wb, ws):
    title(ws, "New store", None)
    put(ws, "A2", '="One new store, month by month for 5 years. Scenario: "&scenario', "note")
    section(ws, 4, "Drivers", 6)
    drivers = [
        ("Store size (sq ft)", "=sqft_new", "link", NUM),
        ("Flagship sales, annual run-rate", "=flag_runrate", "link", USD),
        ("Flagship size (sq ft)", "=sqft_s1", "link", NUM),
        ("Flagship sales per sq ft", "=IF(B7=0,0,B6/B7)", "f", USD),
        ("Sales per sq ft vs flagship", "=prod", "link", PCT0),
        ("Mature annual sales", "=B5*B8*B9", "f", USD),
        ("Gross margin", "=gm_actual+gm_delta", "f", PCT),
        ("Build-out", "=B5*buildout_psf", "f", USD),
        ("Pre-opening costs", "=preopen", "link", USD),
        ("Opening inventory", "=inventory", "link", USD),
        ("Total investment", "=SUM(B12:B14)", "f", USD),
    ]
    for i, (label, formula, kind, fmt) in enumerate(drivers):
        put(ws, f"A{5 + i}", label, "text", bold=label == "Total investment")
        put(ws, f"B{5 + i}", formula, kind, fmt, bold=label == "Total investment")
    name(wb, "ns_mature_sales", ws.title, "B10")
    name(wb, "ns_gm", ws.title, "B11")
    name(wb, "ns_invest", ws.title, "B15")

    # Month-by-month block (rows 46+), written first so the summaries can point at it.
    m0, first, last = 46, 4, 4 + MONTHS - 1  # month 0 in column C, months 1..60 in D..BK
    L = get_column_letter
    months = f"$D${m0}:${L(last)}${m0}"
    years = f"$D${m0 + 1}:${L(last)}${m0 + 1}"
    section(ws, m0 - 1, "Month by month", 6)
    put(ws, f"A{m0}", "Month", "text", bold=True)
    put(ws, f"A{m0 + 1}", "Store year", "text")
    put(ws, f"C{m0}", 0, "in", "0", bold=True)
    put(ws, f"C{m0 + 1}", 0, "in", "0")
    rows = ["Sales ramp, % of mature", "Sales", "Gross profit", "Waste", "Labor", "Rent", "Other operating costs",
            "Store EBITDA (4-wall)", "Opening investment", "Cash flow", "Cumulative cash", "Paid back (month)"]
    R = {label: m0 + 2 + i for i, label in enumerate(rows)}
    for label, r in R.items():
        put(ws, f"A{r}", label, "text" if label != "Paid back (month)" else "note",
            bold=label in ("Store EBITDA (4-wall)", "Cumulative cash"))
    put(ws, f"C{R['Opening investment']}", "=-ns_invest", "f", USD_K)
    put(ws, f"C{R['Cash flow']}", f"=C{R['Store EBITDA (4-wall)']}+C{R['Opening investment']}", "f", USD_K)
    put(ws, f"C{R['Cumulative cash']}", f"=C{R['Cash flow']}", "f", USD_K, bold=True)
    for c in range(first, last + 1):
        x, p = L(c), L(c - 1)
        put(ws, f"{x}{m0}", f"={p}{m0}+1", "f", "0", bold=True)
        put(ws, f"{x}{m0 + 1}", f"=ROUNDUP({x}{m0}/12,0)", "f", "0")
        s = R["Sales"]
        put(ws, f"{x}{R['Sales ramp, % of mature']}",
            f"=IF({x}${m0}>=ramp_months,1,ramp_start+(1-ramp_start)*({x}${m0}-1)/MAX(1,ramp_months-1))", "f", PCT0)
        put(ws, f"{x}{s}", f"=ns_mature_sales/12*{x}{R['Sales ramp, % of mature']}", "f", USD_K)
        put(ws, f"{x}{R['Gross profit']}", f"={x}{s}*ns_gm", "f", USD_K)
        put(ws, f"{x}{R['Waste']}", f"={x}{s}*waste_pct", "f", USD_K)
        put(ws, f"{x}{R['Labor']}", f"=MAX({x}{s}*labor_pct,ns_mature_sales/12*labor_pct*labor_floor)", "f", USD_K)
        put(ws, f"{x}{R['Rent']}", "=sqft_new*rent_psf/12", "f", USD_K)
        put(ws, f"{x}{R['Other operating costs']}", f"={x}{s}*other_pct", "f", USD_K)
        put(ws, f"{x}{R['Store EBITDA (4-wall)']}",
            f"={x}{R['Gross profit']}-{x}{R['Waste']}-{x}{R['Labor']}-{x}{R['Rent']}-{x}{R['Other operating costs']}",
            "f", USD_K, bold=True)
        put(ws, f"{x}{R['Cash flow']}", f"={x}{R['Store EBITDA (4-wall)']}+{x}{R['Opening investment']}", "f", USD_K)
        put(ws, f"{x}{R['Cumulative cash']}", f"={p}{R['Cumulative cash']}+{x}{R['Cash flow']}", "f", USD_K, bold=True)
        put(ws, f"{x}{R['Paid back (month)']}", f"=IF({x}{R['Cumulative cash']}>=0,{x}${m0},999)", "note", NUM)
    for c in range(1, last + 1):
        ws.cell(R["Store EBITDA (4-wall)"], c).border = Border(top=LINE)

    def span(label):
        return f"$D{R[label]}:${L(last)}{R[label]}"

    # Results
    section(ws, 17, "Results", 6)
    payback = f"MIN({span('Paid back (month)')})"
    results = [
        ("Payback (months after opening)", f'=IF({payback}>{MONTHS},"Over {MONTHS}",{payback})', "f", NUM),
        ("Year 3 store EBITDA", "=D32", "f", USD),
        ("Year 3 cash-on-cash return", "=IF(ns_invest=0,0,D32/ns_invest)", "f", PCT),
        ("Sales per sq ft at maturity", "=IF(B5=0,0,ns_mature_sales/B5)", "f", USD),
        ("Store EBITDA margin at maturity",
         "=IF(ns_mature_sales=0,0,ns_gm-waste_pct-labor_pct-other_pct-sqft_new*rent_psf/ns_mature_sales)", "f", PCT),
    ]
    for i, (label, formula, kind, fmt) in enumerate(results):
        put(ws, f"A{18 + i}", label, "text")
        put(ws, f"B{18 + i}", formula, kind, fmt, bold=True, align="right")
    name(wb, "ns_payback", ws.title, "B18")
    name(wb, "ns_y3_ebitda", ws.title, "B19")
    name(wb, "ns_coc", ws.title, "B20")

    # By store year
    section(ws, 24, "By store year", 6)
    for j in range(5):
        c = put(ws, f"{L(2 + j)}25", j + 1, "text", '"Year "0', bold=True, align="right")
        c.border = Border(bottom=LINE)
    yearly = ["Sales", "Gross profit", "Waste", "Labor", "Rent", "Other operating costs", "Store EBITDA (4-wall)"]
    for i, label in enumerate(yearly):
        r = 26 + i
        put(ws, f"A{r}", label, "text", bold=label.startswith("Store EBITDA"))
        for j in range(5):
            x = L(2 + j)
            put(ws, f"{x}{r}", f"=SUMIF({years},{x}$25,{span(label)})", "f", USD, bold=label.startswith("Store EBITDA"))
    put(ws, "A33", "Store EBITDA margin", "text")
    put(ws, "A34", "Cumulative cash at year end", "text")
    for j in range(5):
        x = L(2 + j)
        put(ws, f"{x}33", f"=IF({x}26=0,0,{x}32/{x}26)", "f", PCT)
        put(ws, f"{x}34", f"=INDEX({span('Cumulative cash')},MATCH({x}25*12,{months},0))", "f", USD)
    for c in range(1, 7):
        ws.cell(32, c).border = Border(top=LINE)

    # Sensitivity
    section(ws, 36, "Sensitivity: store EBITDA in a mature year", 6)
    put(ws, "A37", "Rows move sales per sq ft, columns move gross margin, around the live scenario (center).", "note")
    put(ws, "A38", "Sales per sq ft vs flagship  /  gross margin", "text", bold=True)
    for j, d in enumerate((-0.02, -0.01, 0, 0.01, 0.02)):
        c = put(ws, f"{L(2 + j)}38", f"=ns_gm{d:+.2f}" if d else "=ns_gm", "f", PCT, bold=True, align="right")
        c.border = Border(bottom=LINE)
    for i, d in enumerate((-0.2, -0.1, 0, 0.1, 0.2)):
        r = 39 + i
        put(ws, f"A{r}", f"=prod{d:+.1f}" if d else "=prod", "f", PCT0, bold=True, align="right")
        for j in range(5):
            x = L(2 + j)
            put(ws, f"{x}{r}", f"=$B$5*$B$8*$A{r}*({x}$38-waste_pct-labor_pct-other_pct)-$B$5*rent_psf", "f", USD_K)
    ws["D41"].font = Font(name="Arial", size=10, bold=True)
    ws["D41"].border = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)

    widths(ws, {"A": 44, **{L(c): 13 if c <= 6 else 10 for c in range(2, last + 1)}})
    ws.freeze_panes = "B1"
    return dict(month_row=m0, R=R, last=L(last))


def build_plan(wb, ws, act):
    title(ws, "3-Year Plan", None)
    put(ws, "A2", '="Existing stores grow from today\'s run-rate; new stores follow the New Store sheet. '
                  'Scenario: "&scenario', "note")
    put(ws, "A3", "All amounts in $K.", "note")
    headers(ws, 4, ["", "Today (run-rate)", *YEARS, "Notes"])
    cols = ["B", "C", "D", "E"]
    rows = {}

    def line(r, label, formulas, fmt=USD_K, bold=False, kind="f", text=None):
        rows[label] = r
        put(ws, f"A{r}", label, "text", bold=bold)
        for x, f in zip(cols, formulas):
            if f is None:
                continue
            k = kind(x) if callable(kind) else kind
            put(ws, f"{x}{r}", f, k, fmt, bold=bold)
        if text:
            put(ws, f"F{r}", text, "note")

    first_year = lambda x: x == "B"  # noqa: E731
    section(ws, 5, "Drivers", 6)
    line(6, "New stores opened", [None, "=open_2027", "=open_2028", "=open_2029"], NUM, kind="link")
    line(7, "Stores open at year end", [f"=COUNTA(Actuals!$A${act['store_first']}:$A${act['store_total'] - 1})", "=B7+C6", "=C7+D6", "=D7+E6"], NUM)
    line(8, "Selling space (sq ft)", ["=sqft_s1+sqft_s2+sqft_s3", "=B8+C6*sqft_new", "=C8+D6*sqft_new",
                                      "=D8+E6*sqft_new"], NUM)
    line(9, "App ramp, share of 2029 level", [None, "=app_ramp_27", "=app_ramp_28", "=app_ramp_29"], PCT0,
         kind="link")
    line(10, "Active members (estimate)", ["=member_tx/visits", "=B10*(1+g_members)", "=C10*(1+g_members)",
                                           "=D10*(1+g_members)"], NUM,
         text="Today: member transactions in the last 30 days divided by visits per member per month.")
    line(11, "Paying members", [0, "=C10*paid_conv", "=D10*paid_conv", "=E10*paid_conv"], NUM,
         kind=lambda x: "in" if first_year(x) else "f", text="Assumes no paid tier today.")

    section(ws, 13, "Revenue", 6)
    grow = lambda row, rate: [None] + [f"={p}{row}*(1+{rate})" for p in ("B", "C", "D")]  # noqa: E731

    def new_stores(ns_row):
        a = lambda y: f"'New Store'!${'BCD'[y]}${ns_row}"  # noqa: E731
        return [0, f"={a(0)}*C$6", f"={a(1)}*C$6+{a(0)}*D$6", f"={a(2)}*C$6+{a(1)}*D$6+{a(0)}*E$6"]

    line(14, "Grocery, existing stores", ["=grocery_runrate"] + grow(14, "g_grocery")[1:],
         kind=lambda x: "link" if first_year(x) else "f")
    line(15, "Cafe and hot bar, existing stores", ["=cafe_runrate"] + grow(15, "g_cafe")[1:],
         kind=lambda x: "link" if first_year(x) else "f",
         text="Existing stores include Store 3, which is still ramping, so this is conservative.")
    line(16, "New stores", new_stores(26), kind=lambda x: "in" if first_year(x) else "f",
         text="Each year adds that year's openings at year-1 sales and earlier openings at their later-year sales.")
    line(17, "Catering", [0, "=cat_y1", "=C17*(1+g_cat)", "=D17*(1+g_cat)"],
         kind=lambda x: "in" if first_year(x) else ("link" if x == "C" else "f"), text="Not in the data yet.")
    line(18, "App orders (extra sales)", [0] + [f"=({x}14+{x}15+{x}16)*app_pct*{x}9" for x in ("C", "D", "E")],
         kind=lambda x: "in" if first_year(x) else "f")
    line(19, "Membership fees", [f"={x}11*fee" for x in cols])
    line(20, "Total revenue", [f"=SUM({x}14:{x}19)" for x in cols], bold=True)
    line(21, "Growth", [None] + [f"=IF({p}20=0,0,{x}20/{p}20-1)" for p, x in (("B", "C"), ("C", "D"), ("D", "E"))],
         PCT)
    for c in range(1, 6):
        ws.cell(20, c).border = Border(top=LINE)

    section(ws, 23, "Profit", 6)
    line(24, "Existing stores: gross profit",
         ["=B14*grocery_gm+B15*cafe_gm"] + [f"={x}14*(grocery_gm+gm_delta)+{x}15*(cafe_gm+gm_delta)"
                                             for x in ("C", "D", "E")],
         text="Today uses actual margins; later years add the margin change from Inputs.")
    line(25, "Existing stores: waste", ["=(B14+B15)*waste_actual"] + [f"=({x}14+{x}15)*waste_pct" for x in "CDE"])
    line(26, "Existing stores: labor", ["=(B14+B15)*labor_pct_base"] + [f"=({x}14+{x}15)*labor_pct" for x in "CDE"],
         text="Labor, rent and other costs aren't in the data. Today uses the Base assumptions in every scenario.")
    line(27, "Existing stores: rent", ["=(sqft_s1+sqft_s2+sqft_s3)*rent_psf_base"]
         + ["=(sqft_s1+sqft_s2+sqft_s3)*rent_psf"] * 3)
    line(28, "Existing stores: other operating costs", ["=(B14+B15)*other_pct_base"]
         + [f"=({x}14+{x}15)*other_pct" for x in "CDE"])
    line(29, "Existing stores: EBITDA (4-wall)", [f"={x}24-{x}25-{x}26-{x}27-{x}28" for x in cols], bold=True)
    line(30, "New stores: EBITDA (4-wall)", new_stores(32), kind=lambda x: "in" if first_year(x) else "f")
    line(31, "Catering contribution", [f"={x}17*cat_margin" for x in cols])
    line(32, "App orders contribution", [f"={x}18*(gm_actual+gm_delta-app_cost)" for x in cols])
    line(33, "Membership, net of perks", [f"={x}19*(1-perks_cost)" for x in cols])
    line(34, "Head office costs", ["=B20*ho_pct_base"] + [f"={x}20*ho_pct" for x in "CDE"])
    line(35, "Company EBITDA", [f"={x}29+{x}30+{x}31+{x}32+{x}33-{x}34" for x in cols], bold=True)
    line(36, "EBITDA margin", [f"=IF({x}20=0,0,{x}35/{x}20)" for x in cols], PCT)
    for c in range(1, 6):
        ws.cell(35, c).border = Border(top=LINE)

    section(ws, 38, "Cash", 6)
    line(39, "New store investment", [None] + [f"={x}6*ns_invest" for x in "CDE"])
    line(40, "Cash after investment", [None] + [f"={x}35-{x}39" for x in "CDE"], bold=True)
    line(41, "Cumulative, from 2027", [None, "=C40", "=C41+D40", "=D41+E40"], bold=True)
    widths(ws, {"A": 40, "B": 16, "C": 13, "D": 13, "E": 13, "F": 90})
    ws.freeze_panes = "B5"
    return rows


def build_pricing(wb, ws, act):
    title(ws, "Pricing", "Test a price change or a promotion on one category, using its last-30-day price, cost "
                         "and volume.")
    put(ws, "A4", "Category", "text", bold=True)
    put(ws, "B4", DEFAULT_CATEGORY, "in", fill=YELLOW, bold=True)
    ws.merge_cells("B4:C4")
    dv = DataValidation(type="list", formula1="=category_list", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")
    put(ws, "D4", "Pick a category from the list.", "note")
    name(wb, "category", ws.title, "B4")
    look = lambda col: (f"=INDEX(Actuals!${col}${act['cat_first']}:${col}${act['cat_total'] - 1},"  # noqa: E731
                        f"MATCH(category,Actuals!$A${act['cat_first']}:$A${act['cat_total'] - 1},0))")

    section(ws, 6, "Price change", 6)
    rows = [
        (7, "Average shelf price", look("I"), "link", USD2, None),
        (8, "Average unit cost", look("J"), "link", USD2, None),
        (9, "Units per week, all stores", look("K"), "link", NUM, None),
        (10, "Gross margin now", "=IF(B7=0,0,(B7-B8)/B7)", "f", PCT, None),
        (11, "Target margin", look("G"), "link", PCT, "From the Pricing and Margin Policy."),
        (12, "Gap to target", "=B10-B11", "f", PTS, None),
        (13, "Price change", 0.03, "in", '+0.0%;[Red]-0.0%;0.0%', None),
        (14, "Price elasticity", -1.2, "in", DEC,
         "-1.2 means a 1% price rise loses 1.2% of units. Not measured in the demo data; grocery is often "
         "-0.5 to -2.5."),
        (15, "New price", "=B7*(1+B13)", "f", USD2, None),
        (16, "Units per week after", "=MAX(0,B9*(1+B14*B13))", "f", NUM, None),
        (17, "Gross profit per week, now", "=(B7-B8)*B9", "f", USD, None),
        (18, "Gross profit per week, after", "=(B15-B8)*B16", "f", USD, None),
        (19, "Change in gross profit per year", "=(B18-B17)*52", "f", USD, None),
        (20, "New gross margin", "=IF(B15=0,0,(B15-B8)/B15)", "f", PCT, None),
        (21, "Break-even change in units", '=IF(B15-B8<=0,"Below cost",(B7-B8)/(B15-B8)-1)', "f", PTS,
         "The unit change that leaves gross profit flat. Compare it with the expected change."),
    ]
    for r, label, value, kind, fmt, text in rows:
        put(ws, f"A{r}", label, "text", bold=r == 19)
        put(ws, f"B{r}", value, kind, fmt, bold=r == 19, align="right")
        if text:
            put(ws, f"C{r}", text, "note")
    ws["B13"].fill = YELLOW
    ws["B14"].fill = YELLOW

    put(ws, "A23", "Change in gross profit per year: price change (rows) by elasticity (columns)", "text", bold=True)
    elas = (-0.5, -1.0, -1.5, -2.0, -2.5)
    for j, e in enumerate(elas):
        c = put(ws, f"{'BCDEF'[j]}24", e, "in", DEC, bold=True, align="right")
        c.border = Border(bottom=LINE)
    for i, p in enumerate((-0.06, -0.04, -0.02, 0, 0.02, 0.04, 0.06)):
        r = 25 + i
        put(ws, f"A{r}", p, "in", '+0%;-0%;0%', bold=True, align="right")
        for j in range(5):
            x = "BCDEF"[j]
            put(ws, f"{x}{r}", f"=(($B$7*(1+$A{r})-$B$8)*MAX(0,$B$9*(1+{x}$24*$A{r}))-($B$7-$B$8)*$B$9)*52",
                "f", USD)

    section(ws, 33, "Promotion", 6)
    promo = [
        (34, "Discount", 0.20, "in", PCT0, None),
        (35, "Promotion length (weeks)", 2, "in", NUM, None),
        (36, "Expected unit lift", 0.50, "in", PCT0, "Extra units during the promotion vs a normal week."),
        (37, "Share of the discount the vendor pays", 0.50, "in", PCT0, None),
        (38, "Promotion price", "=B7*(1-B34)", "f", USD2, None),
        (39, "Our margin per unit during the promotion", "=B38-B8+B7*B34*B37", "f", USD2, None),
        (40, "Units per week during the promotion", "=B9*(1+B36)", "f", NUM, None),
        (41, "Extra gross profit from the promotion", "=(B39*B40-(B7-B8)*B9)*B35", "f", USD, None),
        (42, "Break-even unit lift", '=IF(B39<=0,"Sells below cost",(B7-B8)/B39-1)', "f", PCT0,
         "The lift the promotion needs to pay for itself."),
        (43, "Markdown policy (Pricing and Margin Policy)",
         '=IF(B34>0.3,"Over the 30% cap: needs COO approval","Within the 30% cap")', "f", None, None),
    ]
    for r, label, value, kind, fmt, text in promo:
        put(ws, f"A{r}", label, "text", bold=r == 41)
        put(ws, f"B{r}", value, kind, fmt, bold=r == 41, align="left" if r == 43 else "right")
        if text:
            put(ws, f"C{r}", text, "note")
    for r in (34, 36):
        ws[f"B{r}"].fill = YELLOW

    put(ws, "A45", "Extra gross profit from the promotion: discount (rows) by unit lift (columns)", "text", bold=True)
    for j, lift in enumerate((0.2, 0.4, 0.6, 0.8, 1.0)):
        c = put(ws, f"{'BCDEF'[j]}46", lift, "in", PCT0, bold=True, align="right")
        c.border = Border(bottom=LINE)
    for i, d in enumerate((0.10, 0.15, 0.20, 0.25, 0.30)):
        r = 47 + i
        put(ws, f"A{r}", d, "in", PCT0, bold=True, align="right")
        for j in range(5):
            x = "BCDEF"[j]
            put(ws, f"{x}{r}", f"=(($B$7*(1-$A{r})-$B$8+$B$7*$A{r}*$B$37)*$B$9*(1+{x}$46)-($B$7-$B$8)*$B$9)*$B$35",
                "f", USD)
    widths(ws, {"A": 44, "B": 13, "C": 13, "D": 13, "E": 13, "F": 13})


def build_checks(wb, ws, act, ns, plan):
    title(ws, "Checks", "Each check works a number out a second way. All should say OK.")
    headers(ws, 4, ["Check", "Value", "Worked out again", "Difference", "Status"])
    st, ct = act["store_total"], act["cat_total"]
    R, last = ns["R"], ns["last"]
    yrs = f"'New Store'!$D${ns['month_row'] + 1}:${last}${ns['month_row'] + 1}"
    sales = f"'New Store'!$D${R['Sales']}:${last}${R['Sales']}"
    ebitda = f"'New Store'!$D${R['Store EBITDA (4-wall)']}:${last}${R['Store EBITDA (4-wall)']}"
    ramp = f"'New Store'!$D${R['Sales ramp, % of mature']}:${last}${R['Sales ramp, % of mature']}"
    plan_ws = quote("3-Year Plan")
    checks = [
        ("Sales: item lines tie to POS transactions", f"=Actuals!C{st}", f"=Actuals!D{st}", USD),
        ("Sales: categories tie to stores", f"=Actuals!C{ct}", f"=Actuals!C{st}", USD),
        ("Cost of goods: categories tie to stores", f"=Actuals!D{ct}", f"=Actuals!E{st}", USD),
        ("Channels add up to total sales", f"=Actuals!B{ct + 5}+Actuals!B{ct + 6}", f"=Actuals!C{ct}", USD),
        ("New store: 60 months tie to the 5 yearly totals (EBITDA)", f"=SUM({ebitda})",
         "=SUM('New Store'!B32:F32)", USD),
        ("New store: sales reach 100% by the ramp month", f"=INDEX({ramp},MIN({MONTHS},MAX(1,ramp_months)))", 1,
         PCT0),
        ("Plan: 2029 new-store sales from the monthly model", f"={plan_ws}!E16",
         f"=open_2027*SUMIF({yrs},3,{sales})+open_2028*SUMIF({yrs},2,{sales})+open_2029*SUMIF({yrs},1,{sales})",
         USD),
        ("Plan: cumulative cash equals the sum of yearly cash", f"={plan_ws}!E41", f"=SUM({plan_ws}!C40:E40)", USD),
        ("Plan: revenue lines add up to total revenue (2029)", f"={plan_ws}!E20",
         f"={plan_ws}!E14+{plan_ws}!E15+{plan_ws}!E16+{plan_ws}!E17+{plan_ws}!E18+{plan_ws}!E19", USD),
        ("Pricing: a 0% price change changes nothing", "=SUM(Pricing!B28:F28)", 0, USD),
    ]
    r0 = 5
    for i, (label, a, b, fmt) in enumerate(checks):
        r = r0 + i
        put(ws, f"A{r}", label, "text")
        put(ws, f"B{r}", a, "link" if "!" in str(a) and not str(a).startswith("=SUM(") else "f", fmt)
        put(ws, f"C{r}", b, "in" if not isinstance(b, str) else ("link" if "!" in b else "f"), fmt)
        put(ws, f"D{r}", f"=B{r}-C{r}", "f", fmt)
        put(ws, f"E{r}", f'=IF(ABS(D{r})<=0.005*MAX(1,ABS(C{r})),"OK","CHECK")', "f", align="center", bold=True)
    r = r0 + len(checks)
    put(ws, f"A{r}", "Inputs: scenario name is valid", "text")
    put(ws, f"B{r}", "=scenario", "link")
    put(ws, f"E{r}", '=IF(ISNUMBER(MATCH(scenario,Inputs!$D$6:$F$6,0)),"OK","CHECK")', "f", align="center", bold=True)
    r += 1
    put(ws, f"A{r}", "Inputs: ramp and labor floor are within 0% to 100%", "text")
    put(ws, f"E{r}", '=IF(AND(ramp_start>0,ramp_start<=1,ramp_months>=1,labor_floor>=0,labor_floor<=1),"OK","CHECK")',
        "f", align="center", bold=True)
    r += 1
    put(ws, f"A{r}", "Pricing: the category is in the Actuals list", "text")
    put(ws, f"B{r}", "=category", "link")
    put(ws, f"E{r}", '=IF(ISNUMBER(MATCH(category,category_list,0)),"OK","CHECK")', "f", align="center", bold=True)
    rng = f"E{r0}:E{r}"
    put(ws, f"A{r + 2}", "Result", "text", bold=True)
    put(ws, f"B{r + 2}", f'=IF(COUNTIF({rng},"OK")=ROWS({rng}),"All "&ROWS({rng})&" checks pass",'
                         f'ROWS({rng})-COUNTIF({rng},"OK")&" of "&ROWS({rng})&" checks need a look")', "f", bold=True)
    name(wb, "checks_status", ws.title, f"B{r + 2}")
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'],
                                                  font=Font(name="Arial", bold=True, color="FF2E845A")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"CHECK"'],
                                                  font=Font(name="Arial", bold=True, color=WHITE),
                                                  fill=PatternFill("solid", fgColor="FFC0392B")))
    widths(ws, {"A": 56, "B": 18, "C": 18, "D": 14, "E": 10})


def build_summary(wb, ws, plan):
    title(ws, "Store planning model", "New store economics, a 3-year plan by channel, and pricing and promotion "
                                      "tests. Built on the Store Copilot demo data, which is synthetic.")
    section(ws, 4, "How to use", 5)
    steps = ["1. Pick a scenario on the Inputs sheet, cell C4: Base, Upside or Downside.",
             "2. Change any blue number. Every black number is a formula and updates.",
             "3. On Pricing, pick a category and test a price change or a promotion.",
             "4. The Checks sheet works key numbers out a second way. It should say all checks pass."]
    for i, s in enumerate(steps):
        put(ws, f"A{5 + i}", s, "text")
    put(ws, "A10", "Blue: an input you can change", "in")
    put(ws, "A11", "Black: a formula", "f")
    put(ws, "A12", "Green: pulled from another sheet", "link")
    put(ws, "A13", "Yellow: the main selectors and levers", "text", fill=YELLOW)

    section(ws, 15, "Key numbers", 5)
    put(ws, "A16", "Scenario", "text", bold=True)
    put(ws, "B16", "=scenario", "link", bold=True, align="right")
    put(ws, "C16", "Checks", "text", bold=True, align="right")
    ws.merge_cells("D16:E16")
    put(ws, "D16", "=checks_status", "link", bold=True, align="right")
    ws.conditional_formatting.add("D16", FormulaRule(formula=['LEFT($D$16,3)="All"'],
                                                     font=Font(name="Arial", bold=True, color="FF2E845A")))
    ws.conditional_formatting.add("D16", FormulaRule(formula=['LEFT($D$16,3)<>"All"'],
                                                     font=Font(name="Arial", bold=True, color="FFC0392B")))
    headers(ws, 18, ["3-year plan ($K)", "Today (run-rate)", *YEARS])
    p = quote("3-Year Plan")
    for i, (label, row, fmt) in enumerate((("Revenue", 20, USD_K), ("Company EBITDA", 35, USD_K),
                                          ("EBITDA margin", 36, PCT), ("Stores open at year end", 7, NUM))):
        r = 19 + i
        put(ws, f"A{r}", label, "text")
        for x in "BCDE":
            put(ws, f"{x}{r}", f"={p}!{x}{row}", "link", fmt)
    put(ws, "A23", "Cash after new-store investment, cumulative", "text")
    for x in "CDE":
        put(ws, f"{x}23", f"={p}!{x}41", "link", USD_K)

    headers(ws, 25, ["One new store", "", "", "", ""])
    ns = [("Total investment", "=ns_invest", USD), ("Mature annual sales", "=ns_mature_sales", USD),
          ("Year 1 sales", "='New Store'!B26", USD), ("Year 3 store EBITDA", "=ns_y3_ebitda", USD),
          ("Payback (months after opening)", "=ns_payback", NUM), ("Year 3 cash-on-cash return", "=ns_coc", PCT)]
    for i, (label, f, fmt) in enumerate(ns):
        put(ws, f"A{26 + i}", label, "text")
        put(ws, f"B{26 + i}", f, "link", fmt, align="right")

    headers(ws, 33, ["Pricing tests", "", "", "", ""])
    put(ws, "A34", '="Changing "&category&" prices by "&TEXT(Pricing!B13,"+0%;-0%")&" changes gross profit by "'
                   '&TEXT(Pricing!B19,"$#,##0;-$#,##0")&" a year (elasticity "&TEXT(Pricing!B14,"0.0")&")."', "f")
    put(ws, "A35", '=IF(ISNUMBER(Pricing!B42),"A "&TEXT(Pricing!B34,"0%")&" promotion on "&category&" needs "'
                   '&TEXT(Pricing!B42,"0%")&" more units to pay for itself. Expected: "&TEXT(Pricing!B36,"0%")&".",'
                   '"A "&TEXT(Pricing!B34,"0%")&" promotion on "&category&" sells below cost.")', "f")

    section(ws, 37, "What's data and what's assumed", 5)
    notes = ["From the demo data (Actuals): sales, cost of goods, margins, waste, basket size, member activity.",
             "Assumed, with a note on Inputs: store sizes, labor, rent, build-out, catering, app orders, "
             "membership fees, head office.",
             "Plan years start in January; new stores open at the start of their year."]
    for i, s in enumerate(notes):
        put(ws, f"A{38 + i}", s, "text")
    widths(ws, {"A": 46, "B": 16, "C": 13, "D": 13, "E": 13})


def build(db: Path = DB_PATH, out: Path = OUT) -> Path:
    act = load_actuals(ensure_db(db))
    wb = Workbook()
    names = ["Summary", "Inputs", "Actuals", "New Store", "3-Year Plan", "Pricing", "Checks"]
    wb.active.title = names[0]
    for n in names[1:]:
        wb.create_sheet(n)
    ws = {n: wb[n] for n in names}
    a = build_actuals(wb, ws["Actuals"], act)
    build_inputs(wb, ws["Inputs"], a)
    ns = build_new_store(wb, ws["New Store"])
    plan = build_plan(wb, ws["3-Year Plan"], a)
    build_pricing(wb, ws["Pricing"], a)
    build_checks(wb, ws["Checks"], a, ns, plan)
    build_summary(wb, ws["Summary"], plan)
    for w in wb.worksheets:
        w.page_setup.orientation = "landscape"
        if w.title != "New Store":  # 60 monthly columns print across pages at full size
            w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
            w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(f"Wrote {build()}")
